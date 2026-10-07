"""Annual district/age survey context. Never changes behavioral detection flags."""
from pathlib import Path
import argparse,hashlib,json,sqlite3,shutil
from contextlib import closing
import pandas as pd
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
VERSION='survey_context_v1'
AGE={2:'20s',3:'30s',4:'40s',5:'50s',6:'60plus'}
SCHEMA='''
CREATE TABLE IF NOT EXISTS ctx_survey_source(
 source_id TEXT PRIMARY KEY, survey TEXT NOT NULL, year INTEGER NOT NULL,
 fingerprint TEXT NOT NULL, available_from TEXT, metadata_json TEXT NOT NULL,
 UNIQUE(survey,year));
CREATE TABLE IF NOT EXISTS ctx_survey_metric(
 source_id TEXT NOT NULL REFERENCES ctx_survey_source(source_id),
 district TEXT NOT NULL, age_band TEXT NOT NULL, metric TEXT NOT NULL,
 value REAL, unit TEXT NOT NULL, n_total INTEGER NOT NULL, n_valid INTEGER NOT NULL,
 n_event INTEGER, n_households INTEGER, effective_n REAL, quality_status TEXT NOT NULL,
 payload_json TEXT NOT NULL, PRIMARY KEY(source_id,district,age_band,metric));
CREATE TABLE IF NOT EXISTS ctx_evidence_guide(
 guide_id TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES ctx_survey_source(source_id),
 age_scope TEXT NOT NULL, topic TEXT NOT NULL, payload_json TEXT NOT NULL);
CREATE VIEW IF NOT EXISTS v_survey_annual AS
SELECT s.survey,s.year,s.available_from,m.* FROM ctx_survey_metric m
JOIN ctx_survey_source s USING(source_id);
'''

def dumps(x):return json.dumps(x,ensure_ascii=False,allow_nan=False,default=str)
def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def read_book(path):
 w=load_workbook(path,read_only=True,data_only=True)
 try:
  s=w.worksheets[0];it=s.values;headers=next(it)
  if len(set(headers))!=len(headers):raise ValueError('Duplicate data headers')
  guide={}
  for r in w.worksheets[1].iter_rows(values_only=True):
   if r[0]:guide[str(r[0]).upper()]=list(r)
  return w,it,headers,guide
 except Exception:w.close();raise

def normalize(d,mapping):
 d=d.copy()
 cols={k:mapping[k] for k in ['parent','sick','money','emotion','family','outside']}
 for name,col in cols.items():
  v=pd.to_numeric(d[col],errors='coerce')
  allowed=[1,2] if name in ['parent','sick','money','emotion'] else [1,2,3,4,5]
  if (~v.isna() & ~v.isin(allowed)).any():raise ValueError(f'Unexpected response code: {col}')
  d[name]=v
 for name in ['sick','money','emotion']:
  if (d.parent.eq(2)&d[name].eq(1)).any():raise ValueError('Contradictory support branch')
  d[name+'_structural']=d.parent.eq(2)&d[name].isna()
  d.loc[d[name+'_structural'],name]=2
 return d

def metric_row(group,source,district,age,metric,values,unit,definition,event=None):
 weight=pd.to_numeric(group.weight,errors='coerce')
 if weight.isna().any() or (~np.isfinite(weight)).any() or weight.le(0).any():raise ValueError('Invalid survey weight')
 valid=values.notna();v=values[valid];w=weight[valid]
 n=len(v);neff=float(w.sum()**2/(w*w).sum()) if n else None
 value=float((v*w).sum()/w.sum()) if n else None
 events=int(event[valid].sum()) if event is not None else None
 status='available' if n>=30 and neff>=30 else 'small_sample' if n else 'unavailable'
 notes=[]
 if events is not None and (events<5 or n-events<5):notes.append('sparse_response_count')
 if n and len(group)>n:notes.append('item_nonresponse_excluded')
 payload={'definition':definition,'notes':notes,'weight':'wtb1' if source.startswith('seoul_') else 'source_specific',
 'uncertainty':'Effective n reflects weight dispersion only. Household clustering/strata are not corrected. No significance or confidence claim.',
 'structural_no_support_counts':{k:int(group[k+'_structural'].sum()) for k in ['sick','money','emotion'] if k+'_structural' in group}}
 households=int(group.loc[valid,'household_id'].nunique()) if 'household_id' in group else None
 return (source,district,age,metric,value,unit,len(group),n,events,households,neff,status,dumps(payload))

def summarize_seoul(d,source,mapping):
 d=normalize(d,mapping);d['weight']=d[mapping['weight']]
 d['household_id']=d[mapping['household']]
 rows=[]
 for district,df in [('gangnam',d[d[mapping['district']].eq(680)]),('seoul',d)]:
  for age in ['all15plus',*AGE.values()]:
   g=df if age=='all15plus' else df[df[mapping['age']].eq(next(k for k,v in AGE.items() if v==age))]
   for k in ['parent','sick','money','emotion']:
    values=g[k].eq(2).astype(float).where(g[k].isin([1,2]))
    rows.append(metric_row(g,source,district,age,'no_support_'+k,values,'proportion','No helper in this situation; general-no branch inferred for the three specific support items.',values))
   complete=g[['sick','money','emotion']].notna().all(axis=1)
   none=g[['sick','money','emotion']].eq(2).all(axis=1).astype(float).where(complete)
   rows.append(metric_row(g,source,district,age,'no_support_all3',none,'proportion','No helper for all three situations (illness, money, emotional conversation).',none))
   for k in ['family','outside']:
    values=(g[k]-1)/4*10
    rows.append(metric_row(g,source,district,age,'loneliness_'+k+'_mean10',values,'score_0_10','Weighted mean of source five-category loneliness item, rescaled (x-1)/4*10.'))
    event=g[k].ge(4).astype(float).where(g[k].notna())
    rows.append(metric_row(g,source,district,age,'loneliness_'+k+'_agree',event,'proportion','Responses 4 or 5 to the specific loneliness item.',event))
   joint=(none.eq(1)&g.outside.ge(4)).astype(float).where(none.notna()&g.outside.notna())
   rows.append(metric_row(g,source,district,age,'no_support3_and_lonely_outside',joint,'proportion','Intersection of no helper in all three situations and outside-family loneliness response 4/5. Not an isolation diagnosis.',joint))
   for k,col in mapping.get('extra_loneliness',{}).items():
    v=pd.to_numeric(g[col],errors='coerce')
    if (~v.isna()&~v.isin([1,2,3,4,5])).any():raise ValueError('Unexpected new loneliness code')
    rows.append(metric_row(g,source,district,age,k+'_mean10',(v-1)/4*10,'score_0_10','New 2025 item, no preceding-year comparison.'))
    event=v.ge(4).astype(float).where(v.notna())
    rows.append(metric_row(g,source,district,age,k+'_agree',event,'proportion','New 2025 item responses 4/5, no preceding-year comparison.',event))
 return rows

def packet_seoul(year,option):
 path=Path(option['path'])
 if str(year) not in path.name:raise ValueError('Configured year does not match source filename')
 fingerprint=sha(path);w,it,h,guide=read_book(path)
 mapping=option['mapping'];required=list(dict.fromkeys([mapping[k] for k in ['district','age','household','weight','parent','sick','money','emotion','family','outside']]+list(mapping.get('extra_loneliness',{}).values())))
 try:
  for c in required:
   if c not in h:raise ValueError(f'{year}: missing required column {c}')
  # Check semantics, not just variable numbers, before accepting a new annual file.
  for k,term in [('parent','도움'),('sick','아플'),('money','금전'),('emotion','이야기'),('family','외로움'),('outside','외로움')]:
   if term not in str(guide.get(mapping[k].upper())):raise ValueError(f'{year}: unexpected codebook meaning for {k}')
  if '20대' not in str(guide.get(mapping['age'].upper())) or '60세 이상' not in str(guide.get(mapping['age'].upper())):raise ValueError('Unverified age scheme')
  idx=[h.index(c) for c in required];data=pd.DataFrame([[r[j] for j in idx] for r in it],columns=required)
  if not data[mapping['age']].isin([1,2,3,4,5,6]).all():raise ValueError('Unverified or missing age category')
 finally:w.close()
 source=f'seoul_{year}'
 metadata={'path':str(path),'sha256':fingerprint,'rows':len(data),'gangnam_rows':int(data[mapping['district']].eq(680).sum()),'mapping':mapping,'version':VERSION,'geography':'district','age_scheme':'published DEW2','branch_rule':'parent=2: missing sick/money/emotion becomes structural absence','codebook':{c:guide.get(c.upper()) for c in required},'survey_design':'repeated cross-section, not a person panel','response_wording_break':year>=2024}
 return {'source_id':source,'survey':'seoul_survey','year':year,'fingerprint':fingerprint,'available_from':option.get('available_from'),'metadata':metadata,'rows':summarize_seoul(data,source,mapping),'guides':[]}

def packet_youth(kind,option):
 path=Path(option['path']);fingerprint=sha(path);w,it,h,guide=read_book(path)
 fields=['SEQ','SQ4','HIKI_C','SQ11_1_RR','WT_ALL3','WT_ALL2'] if kind=='household' else ['SEQ','SQ4','KEY_1','SQ2','QCL_3']
 try:
  idx=[h.index(c) for c in fields];d=pd.DataFrame([[r[j] for j in idx] for r in it],columns=fields)
 finally:w.close()
 source='youth2022_'+kind;label='HIKI_C' if kind=='household' else 'KEY_1'
 if kind=='household':
  d['age_band']=d.SQ11_1_RR.map({1:'19to24',2:'25to29',3:'30to34',4:'35to39'})
  d['weight']=d.WT_ALL3;allowed=[1,2,3]
 else:
  ages=pd.to_numeric(d.SQ2,errors='coerce')
  if ages.isna().any() or (~ages.between(19,39)).any():raise ValueError('Youth age outside verified range')
  d['age_band']=pd.cut(ages,[18,24,29,34,39],labels=['19to24','25to29','30to34','35to39']).astype(str)
  d['weight']=1.;allowed=[1,2]
 if d.age_band.isna().any() or (~d[label].isin(allowed)).any():raise ValueError('Invalid youth codes')
 d['household_id']=d.SEQ;rows=[]
 for district,df in [('gangnam',d[d.SQ4.eq(1)]),('seoul',d)]:
  for age in ['19to39','19to24','25to29','30to34','35to39']:
   g=df if age=='19to39' else df[df.age_band.eq(age)]
   values=(g[label].isin([1,2]) if kind=='household' else g[label].eq(1)).astype(float)
   rows.append(metric_row(g,source,district,age,'survey_isolated_or_withdrawn',values,'proportion','2022 household weighted classification' if kind=='household' else '2022 youth respondent unweighted classification share, not population prevalence',values))
 metadata={'path':str(path),'sha256':fingerprint,'rows':len(d),'unique_households_or_respondents':int(d.SEQ.nunique()),'gangnam_rows':int(d.SQ4.eq(1).sum()),'weight':'WT_ALL3' if kind=='household' else 'none: unweighted respondent composition','scope':'19-39 years; 19to24 is not DB1 20s','static_reference':True,'version':VERSION}
 if kind=='household':metadata['weight_audit']={'WT_ALL2_sum':float(d.WT_ALL2.sum()),'WT_ALL3_sum':float(d.WT_ALL3.sum()),'codebook_WT_ALL2':guide.get('WT_ALL2'),'note':'WT_ALL2 population sum and guide population text differ. Do not report expanded population counts without provider clarification.'}
 texts=[('movement','외출 빈도와 지속기간은 이동 위축의 해석 근거이다. 이동 데이터와 설문 응답자는 연결되지 않으며 정확도 검증 자료가 아니다.', ['A4x1','A5x1'] if kind=='household' else ['A7','A8']),('communication','대면 교류·지원망을 구분한다. 통화·문자 감소가 모든 소통의 감소를 뜻하지 않는다.', ['A6x1','A7x1','A8x1'] if kind=='household' else ['A12','A13','A14','A17'])]
 if kind=='youth':texts += [('consumption','청년조사 A10은 물건 구매·식당 주문의 일시적 만남을 제외한다. 소비 감소는 보조 근거이며 사회적 접촉 수로 바꾸지 않는다.',['A10']),('services','경제·정서·활동·관계 지원 필요를 참고한다. 실제 지원기관과 이용 조건은 별도 최신 자료로 확인한다.',['C4_1 TO C4_9'])]
 guides=[(source+'_'+topic,source,'19to39',topic,dumps({'text':text,'source_items':fields,'survey_year':2022,'release_date':'2024-05-22','limitations':'Static regional reference. Do not validate 2025/2026 dong detection or increase age risk scores with this dataset.'})) for topic,text,fields in texts]
 return {'source_id':source,'survey':'seoul_youth_'+kind,'year':2022,'fingerprint':fingerprint,'available_from':'2024-05-22','metadata':metadata,'rows':rows,'guides':guides}

def refresh_views(con):
 con.executescript('''
 DROP VIEW IF EXISTS v_survey_age_trend;
 CREATE VIEW v_survey_age_trend AS
 SELECT a.*,b.value previous_value,a.value-b.value change_value,
 CASE WHEN b.source_id IS NULL THEN 'no_previous_year'
 WHEN a.year=2024 AND a.metric LIKE 'loneliness_%' THEN 'response_wording_changed'
 WHEN a.quality_status!='available' OR b.quality_status!='available' THEN 'small_sample_review'
 ELSE 'descriptive_only_not_significance' END comparison_status
 FROM v_survey_annual a LEFT JOIN v_survey_annual b
 ON b.survey=a.survey AND b.year=a.year-1 AND b.district=a.district
 AND b.age_band=a.age_band AND b.metric=a.metric;
 ''')
 if not con.execute("SELECT 1 FROM sqlite_master WHERE name='v_a123_age_elder_context'").fetchone():return
 for name,available in [('v_a123_age_survey_context',False),('v_a123_age_survey_available_context',True)]:
  timing="AND s.available_from IS NOT NULL AND s.available_from<=date(substr(b.date,1,7)||'-01','+1 month','-1 day')" if available else ''
  con.executescript(f'''DROP VIEW IF EXISTS {name}; CREATE VIEW {name} AS
  SELECT b.*,s.year survey_year,s.available_from survey_available_from,
  'district_not_dong' survey_geography_alignment,'reference_not_person_link' survey_evidence_role,
  CASE WHEN s.year IS NULL THEN 'unavailable' WHEN s.year=CAST(substr(b.date,1,4) AS INTEGER) THEN 'same_year_retrospective' ELSE 'prior_year_reference' END survey_time_alignment,
  CASE WHEN s.source_id IS NULL THEN NULL ELSE (SELECT json_group_object(m.metric,json_object('value',m.value,'unit',m.unit,'n_valid',m.n_valid,'n_event',m.n_event,'effective_n',m.effective_n,'quality_status',m.quality_status,'metadata',json(m.payload_json))) FROM ctx_survey_metric m WHERE m.source_id=s.source_id AND m.district='gangnam' AND m.age_band=b.age_band) END survey_context_json
  FROM v_a123_age_elder_context b LEFT JOIN ctx_survey_source s ON s.source_id=(
   SELECT s.source_id FROM ctx_survey_source s WHERE s.survey='seoul_survey'
   AND s.year<=CAST(substr(b.date,1,4) AS INTEGER) {timing} ORDER BY s.year DESC LIMIT 1);
  ''')

def export(db,output):
 output=Path(output);output.mkdir(parents=True,exist_ok=True)
 with closing(sqlite3.connect(db)) as c:
  if not c.execute("SELECT 1 FROM sqlite_master WHERE name='ctx_survey_source'").fetchone():return {'status':'not_initialized'}
  refresh_views(c)
  for name in ['v_survey_annual','v_survey_age_trend','ctx_evidence_guide','v_a123_age_survey_context','v_a123_age_survey_available_context']:
   if not c.execute('SELECT 1 FROM sqlite_master WHERE name=?',(name,)).fetchone():continue
   d=pd.read_sql_query('SELECT * FROM '+name,c)
   tmp=output/(name+'.csv.tmp');d.to_csv(tmp,index=False,encoding='utf-8-sig');tmp.replace(output/(name+'.csv'))
  report={'sources':c.execute('SELECT survey,year FROM ctx_survey_source ORDER BY survey,year').fetchall(),'metric_rows':c.execute('SELECT COUNT(*) FROM ctx_survey_metric').fetchone()[0],'guide_rows':c.execute('SELECT COUNT(*) FROM ctx_evidence_guide').fetchone()[0],'detection_scores_modified':False}
  for name in ['v_a123_age_survey_context','v_a123_age_survey_available_context']:
   if c.execute('SELECT 1 FROM sqlite_master WHERE name=?',(name,)).fetchone():report[name]={'rows':c.execute('SELECT COUNT(*) FROM '+name).fetchone()[0],'linked':c.execute('SELECT COUNT(*) FROM '+name+' WHERE survey_year IS NOT NULL').fetchone()[0]}
 (output/'survey_context_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');return report

def run(db,settings,output):
 options=settings.get('survey_context',{});packets=[];unchanged=[];publication_updates=[]
 with closing(sqlite3.connect(db)) as c:
  c.execute('PRAGMA foreign_keys=ON');c.executescript(SCHEMA)
  candidates=[('seoul_survey',int(year),option,lambda y=int(year),o=option:packet_seoul(y,o)) for year,option in options.get('seoul_years',{}).items()]
  candidates += [('seoul_youth_'+kind,2022,option,lambda k=kind,o=option:packet_youth(k,o)) for kind,option in options.get('youth_reference',{}).items()]
  for survey,year,option,reader in candidates:
   if year<2022:raise ValueError('Unexpected reference year')
   path=Path(option['path']);fingerprint=sha(path)
   date=option.get('available_from') if survey=='seoul_survey' else '2024-05-22'
   if date and (pd.Timestamp(date).strftime('%Y-%m-%d')!=date or date<f'{year}-01-01'):raise ValueError('Invalid availability date')
   stored=c.execute('SELECT fingerprint,metadata_json,available_from FROM ctx_survey_source WHERE survey=? AND year=?',(survey,year)).fetchone()
   if stored:
    if stored[0]!=fingerprint:raise ValueError(f'{survey}/{year}: source correction requires explicit versioned migration')
    prior=json.loads(stored[1])
    if option.get('mapping') is not None and prior.get('mapping')!=option['mapping']:raise ValueError('Historical mapping changed; explicit migration required')
    if stored[2]!=date:
     if stored[2] is None and date:publication_updates.append((date,survey,year))
     else:raise ValueError('Publication metadata change requires explicit migration')
    unchanged.append([survey,year]);continue
   p=reader()
   if fingerprint!=sha(path):raise ValueError('Source changed while reading')
   if p['available_from']:
    date=pd.Timestamp(p['available_from']).strftime('%Y-%m-%d')
    if date!=p['available_from'] or date<f'{year}-01-01':raise ValueError('Invalid availability date')
   packets.append(p)
  # Parse and validate all sources before saving any metric.
  with c:
   for update in publication_updates:c.execute('UPDATE ctx_survey_source SET available_from=? WHERE survey=? AND year=?',update)
   for p in packets:
    archive=ROOT/'Context/data/raw'/p['fingerprint'];archive.mkdir(parents=True,exist_ok=True)
    source=Path(p['metadata']['path']);target=archive/source.name
    if target.exists() and sha(target)!=p['fingerprint']:raise ValueError('Archive checksum mismatch')
    if not target.exists():shutil.copy2(source,target)
    p['metadata']['archive']=str(target.relative_to(ROOT))
    c.execute('INSERT INTO ctx_survey_source VALUES(?,?,?,?,?,?)',(p['source_id'],p['survey'],p['year'],p['fingerprint'],p['available_from'],dumps(p['metadata'])))
    c.executemany('INSERT INTO ctx_survey_metric VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',p['rows'])
    c.executemany('INSERT INTO ctx_evidence_guide VALUES(?,?,?,?,?)',p['guides'])
  refresh_views(c)
 return {**export(db,output),'inserted_sources':len(packets),'unchanged_sources':unchanged,'publication_dates_filled':len(publication_updates)}

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--config',default=str(ROOT/'config/db1_config.json'));p.add_argument('--db');p.add_argument('--output');p.add_argument('--export-only',action='store_true');a=p.parse_args()
 settings=json.loads(Path(a.config).read_text(encoding='utf-8'));db=Path(a.db) if a.db else ROOT/settings['database'];output=Path(a.output) if a.output else ROOT/settings.get('exports','outputs/integrated')
 print(dumps(export(db,output) if a.export_only else run(db,settings,output)))
