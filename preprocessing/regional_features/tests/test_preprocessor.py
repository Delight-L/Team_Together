import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
import stat
import contextlib
from unittest.mock import patch
import openpyxl
from preprocessing.archives import discover
from preprocessing.readers import read_workbook, TELECOM, INTEREST, number
from preprocessing.features import assemble, months
from preprocessing.temporary import temporary_directory
from preprocessing.validation import compare
from main import main

BASE = Path(__file__).resolve().parents[1]
LIMITS = json.loads((BASE / 'config.json').read_text())['archive_limits']


def manifest():
    return {'files': [], 'name_normalizations': set()}


class ArchiveTests(unittest.TestCase):
    def test_nested_and_duplicate(self):
        with temporary_directory() as d:
            root = Path(d)
            inner = io.BytesIO()
            with zipfile.ZipFile(inner, 'w') as z:
                z.writestr('2022.1월_통신.xlsx', b'fixture')
                z.writestr('2022.1월_copy.xlsx', b'fixture')
            archive = root / 'raw.zip'
            with zipfile.ZipFile(archive, 'w') as z:
                z.writestr('nested.zip', inner.getvalue())
            log = manifest()
            files = discover([archive], root / 'scratch', LIMITS, log)
            self.assertEqual(len(files), 1)
            self.assertTrue(any(e['status'] == 'duplicate' for e in log['files']))

    def test_month_aware_identity_and_loose_limits(self):
        with temporary_directory() as d:
            root=Path(d)
            archive=root/'input.zip'
            with zipfile.ZipFile(archive,'w') as z:
                z.writestr('notes.txt',b'same bytes')
                z.writestr('2022.1월.xlsx',b'same bytes')
                z.writestr('2022.2월.xlsx',b'same bytes')
            self.assertEqual(len(discover([archive],root/'scratch',LIMITS,manifest())),2)
            loose=root/'2022.1월.xlsx';loose.write_bytes(b'123')
            with self.assertRaises(ValueError):discover([loose],root/'scratch',dict(LIMITS,max_file_bytes=1),manifest())
            with self.assertRaises(ValueError):discover([loose],root/'scratch',dict(LIMITS,max_files=0),manifest())

    def test_explicit_symlink_checked_before_resolve(self):
        with patch('pathlib.Path.is_symlink', return_value=True), patch('pathlib.Path.resolve', side_effect=AssertionError('resolve called first')):
            with self.assertRaisesRegex(ValueError,'심볼릭'):
                discover(['link.zip'],BASE/'unused',LIMITS,manifest())

    def test_unsafe_and_corrupt_and_limits(self):
        for member in ['../escape.xlsx', 'C:/escape.xlsx', '/escape.xlsx']:
            with self.subTest(member=member), temporary_directory() as d:
                p = Path(d) / 'bad.zip'
                with zipfile.ZipFile(p, 'w') as z:
                    z.writestr(member, b'x')
                with self.assertRaises(ValueError):
                    discover([p], Path(d)/'scratch', LIMITS, manifest())
        with temporary_directory() as d:
            p = Path(d)/'bad.zip'
            p.write_bytes(b'corrupt')
            with self.assertRaises(zipfile.BadZipFile):
                discover([p], Path(d)/'scratch', LIMITS, manifest())
            with zipfile.ZipFile(p,'w') as z:
                z.writestr('a.xlsx',b'12345')
            for override in [{'max_files':0}, {'max_depth':0}, {'max_total_bytes':1}, {'max_file_bytes':1}]:
                with self.assertRaises(ValueError):
                    discover([p],Path(d)/'scratch',dict(LIMITS,**override),manifest())

    def test_symlink(self):
        with temporary_directory() as d:
            p=Path(d)/'link.zip'
            info=zipfile.ZipInfo('link.xlsx')
            info.create_system=3
            info.external_attr=(stat.S_IFLNK | 0o777) << 16
            with zipfile.ZipFile(p,'w') as z:z.writestr(info,'target')
            with self.assertRaises(ValueError):discover([p],Path(d)/'scratch',LIMITS,manifest())


class FeatureTests(unittest.TestCase):
    def fixture(self, path, kind, negative=False):
        book = openpyxl.Workbook()
        sheet = book.active
        headers = ['행정동코드','자치구','행정동명','성별','연령대']
        if kind == 'telecom':
            headers += ['총인구'] + list(dict.fromkeys(x for v in TELECOM.values() for x in v[:2] if x))
        else:
            headers += ['1인가구수'] + list(INTEREST.values())
        sheet.append(headers)
        for age,pop,value in [(20,10,2),(30,30,6)]:
            row = dict.fromkeys(headers, value)
            row.update(행정동코드=1123074,자치구='강남구',행정동명='개포3동',성별=1,연령대=age,총인구=pop)
            for _,excluded,_ in TELECOM.values():
                if excluded:
                    row[excluded] = pop+1 if negative else pop/2
            if kind == 'interest':
                row['1인가구수']=pop
                for field in INTEREST.values():
                    row[field]=0
            sheet.append([row[k] for k in headers])
        book.save(path)

    def test_weighted_aggregation_and_structural_flag(self):
        with temporary_directory() as d:
            data={}; areas={'1123074':'일원2동'}; log=manifest()
            for kind in ['telecom','interest']:
                path=Path(d)/f'{kind}.xlsx';self.fixture(path,kind)
                month,k,groups,_,count=read_workbook(path,'2024.1월_data.xlsx',areas,log,{'sex':[1],'age':[20,30]})
                data[(month,k)]=groups
                self.assertEqual(count,2)
            rows=assemble(data,['2024-01-01'],areas,{'2024-01-01':dict(rain_days=1,rainfall_mm=2,snow_days=0)})
            self.assertEqual(rows[0]['call_contacts'],5)
            self.assertEqual(rows[0]['weekday_move_count'],5)
            self.assertEqual(rows[0]['weekday_count_est_ratio'],0.5)
            self.assertEqual(rows[0]['comm_low_rate'],0)
            self.assertTrue(rows[0]['interest_structural_issue'])
            self.assertEqual(rows[0]['weekday_days']+rows[0]['weekend_days'],31)
            self.assertTrue(log['name_normalizations'])
            with self.assertRaisesRegex(ValueError, '인구 셀 누락'):
                read_workbook(path,'2024.1월_data.xlsx',areas,log)
            with self.assertRaises(ValueError):
                assemble(data,['2024-02-01'],areas,{})
            data[('2024-01-01','interest')]['1123074']['households']=0
            with self.assertRaisesRegex(ValueError,'분모 0'):
                assemble(data,['2024-01-01'],areas,{'2024-01-01':dict(rain_days=0,rainfall_mm=0,snow_days=0)})

    def test_missing_column_and_normalized_content_identity(self):
        with temporary_directory() as d:
            p=Path(d)/'fixture.xlsx'; self.fixture(p,'interest')
            args=('2024.1월_data.xlsx',{'1123074':'일원2동'},manifest(),{'sex':[1],'age':[20,30]})
            original=read_workbook(p,*args)[3]
            book=openpyxl.load_workbook(p);sheet=book.active
            sheet.cell(2,7,1)
            book.save(p);book.close()
            self.assertNotEqual(original,read_workbook(p,*args)[3])
            book=openpyxl.load_workbook(p);sheet=book.active
            sheet.cell(1,7,'missing');book.save(p);book.close()
            with self.assertRaises(ValueError):read_workbook(p,*args)

    def test_bad_numeric_and_negative_weight(self):
        for value in [None,'',float('nan'),float('inf'),-1,True,'unknown']:
            with self.assertRaises(ValueError):number(value,'test')
        with temporary_directory() as d:
            p=Path(d)/'fixture.xlsx';self.fixture(p,'telecom',negative=True)
            with self.assertRaises(ValueError):
                read_workbook(p,'2024.1월_data.xlsx',{'1123074':'일원2동'},manifest())

    def test_calendar(self):
        self.assertEqual(len(months('2022-01','2025-12')),48)
        with self.assertRaises(ValueError):months('2025-12','2022-01')


class CliTests(unittest.TestCase):
    def test_reproducible_custom_weather_failure_and_malformed_config(self):
        with temporary_directory() as d:
            root=Path(d);input_dir=root/'input';input_dir.mkdir();out=root/'output'
            areas=json.loads((BASE/'reference/areas.json').read_text(encoding='utf-8'))
            for kind in ['telecom','interest']:
                book=openpyxl.Workbook();sheet=book.active
                headers=['행정동코드','자치구','행정동명','성별','연령대']
                headers += ['총인구'] + list(dict.fromkeys(x for v in TELECOM.values() for x in v[:2] if x)) if kind=='telecom' else ['1인가구수']+list(INTEREST.values())
                sheet.append(headers)
                for code,name in areas.items():
                    for sex in [1,2]:
                        for age in range(20,80,5):
                            row=dict.fromkeys(headers,2)
                            row.update(행정동코드=int(code),자치구='강남구',행정동명=name,성별=sex,연령대=age,총인구=100,**{'1인가구수':20})
                            sheet.append([row[k] for k in headers])
                book.save(input_dir/f'2024.1월_{kind}.xlsx');book.close()
            # Outside-period workbook must be excluded before schema validation.
            book=openpyxl.Workbook();book.active.append(['unsupported']);book.save(input_dir/'2023.1월_old.xlsx');book.close()
            weather=root/'weather.csv';weather.write_text('date,rain_days,rainfall_mm,snow_days\n2024-01-01,1,2,0\n',encoding='utf-8')
            cfg=json.loads((BASE/'config.json').read_text(encoding='utf-8'))
            cfg.update(start_month='2024-01',end_month='2024-01',input=str(input_dir),output=str(out),weather=str(weather))
            path=root/'config.json';path.write_text(json.dumps(cfg),encoding='utf-8')
            args=['--full','--config',str(path),'--output',str(out)]
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main(args),0)
                result=out/'gangnam_analysis2_feature_table_2022_2025.csv';first=result.read_bytes()
                self.assertEqual(main(args),0)
                self.assertEqual(first,result.read_bytes())
                log=json.loads((out/'preprocessing_manifest.json').read_text(encoding='utf-8'))
                self.assertEqual(log['weather']['provenance']['source'],'사용자 지정 날씨 CSV')
                self.assertTrue(any(f['status']=='excluded_outside_period' for f in log['files']))
                # Cleanup failure leaves previous published CSV untouched.
                @contextlib.contextmanager
                def failed_cleanup(parent):
                    with temporary_directory(parent) as folder:
                        yield folder
                    raise OSError('simulated cleanup failure')
                with patch('main.temporary_directory',failed_cleanup):
                    self.assertEqual(main(args),1)
                self.assertEqual(first,result.read_bytes())
                (input_dir/'2024.1월_interest.xlsx').unlink()
                self.assertEqual(main(args),1)
                report=json.loads((out/'validation_report.json').read_text(encoding='utf-8'))
                self.assertTrue(report['existing_output_is_not_current_run'])
                self.assertEqual(first,result.read_bytes())
                path.write_text('{ malformed',encoding='utf-8')
                self.assertEqual(main(args),1)
                self.assertEqual(json.loads((out/'validation_report.json').read_text(encoding='utf-8'))['status'],'FAIL')
                for tolerance in [-1,float('inf'),float('nan')]:
                    cfg['comparison_absolute_tolerance']=tolerance
                    path.write_text(json.dumps(cfg),encoding='utf-8')
                    self.assertEqual(main(args),1)
                    with self.assertRaises(ValueError):compare([],root/'nonexistent',[],tolerance)


if __name__ == '__main__':
    unittest.main()
