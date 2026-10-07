"""검증입력 엑셀의 확인된 2025년 동·월 실적을 대시보드 JSON으로 가져온다.

실행: python -m legacy.streamlit_dashboard.tools.import_validation <2025_동별_검증입력_청구대응표.xlsx>
"""
import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from legacy.streamlit_dashboard.tools.detection_data import METRICS
from legacy.streamlit_dashboard.tools.validation_data import ACTUAL_FIELDS, VALIDATION_FILE


HEADERS = (
    "기준연월", "행정동코드", "행정동", "탐지 단계", "조사 대상 가구", "조사 완료 가구",
    "신규 발굴 가구", "고위험 가구", "중위험 가구", "저위험 가구", "연락 시도",
    "상담 완료", "방문 확인", "지원 필요", "서비스 연계", "사후관리",
    "실적 문서번호", "실적 기준기간", "단위(명/가구)", "중복 처리 기준", "비고",
)


def import_workbook(source: Path, output: Path = VALIDATION_FILE) -> int:
    known = {
        (r["month"], str(r["dong_code"])): r["dong"]
        for r in json.loads(METRICS.read_text(encoding="utf-8"))["rows"]
        if r["month"].startswith("2025-") and r["dong"] != "개포3동"
    }
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book["검증입력"]
    rows = sheet.iter_rows(values_only=True)
    for row in rows:
        if tuple(row[:len(HEADERS)]) == HEADERS:
            break
    else:
        raise ValueError("검증입력 시트의 열 제목을 찾지 못했습니다.")

    records = []
    seen = set()
    for row in rows:
        if not row[0]:
            continue
        month, code, dong = str(row[0]).strip(), str(row[1]).strip(), str(row[2]).strip()
        key = (month, code)
        if key not in known or known[key] != dong:
            if dong == "개포3동" and all(v is None for v in row[4:16]):
                continue
            raise ValueError(f"탐지표에 없는 2025년 동·월: {month} {code} {dong}")
        if key in seen:
            raise ValueError(f"동·월 중복: {month} {dong}")
        seen.add(key)
        values = row[4:16]
        if all(v is None for v in values):
            continue
        metadata = row[16:20]
        if not all(v is not None and str(v).strip() for v in metadata):
            raise ValueError(f"실적 출처·기준기간·단위·중복 처리 기준 누락: {month} {dong}")
        if str(metadata[1]).strip() != month:
            raise ValueError(f"기준기간이 탐지 월과 다릅니다. 넓은 기간의 실적을 월별로 배분하지 마세요: {month} {dong}")
        cleaned = {}
        for field, value in zip(ACTUAL_FIELDS, values):
            if value is None:
                cleaned[field] = None
            elif isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or int(value) != value:
                raise ValueError(f"실적 값이 0 이상의 정수가 아닙니다: {month} {dong} {field}")
            else:
                cleaned[field] = int(value)
        records.append({
            "month": month, "dongCode": code, "dong": dong, **cleaned,
            "documentNumber": str(metadata[0]).strip(), "period": str(metadata[1]).strip(),
            "unit": str(metadata[2]).strip(), "dedupRule": str(metadata[3]).strip(),
            "note": str(row[20]).strip() if row[20] is not None else "",
        })
    book.close()
    result = {"source": source.name, "asOf": date.today().isoformat(), "records": records}
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return len(records)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    count = import_workbook(args.workbook)
    print(f"2025년 동·월 실적 {count}건 반영. 빈칸은 0으로 처리하지 않았습니다.")
