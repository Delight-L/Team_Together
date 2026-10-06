"""편집 가능한 지역 변화 검토 보고서 초안을 만듭니다."""

from io import BytesIO


def build_report(context, rows, author, department, opinion, source_note, workflow=None):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from datetime import datetime
    from zoneinfo import ZoneInfo

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(11 if name == "Normal" else 15 if name == "Title" else 12)
        style.paragraph_format.space_after = Pt(7)
    document.add_paragraph(("[시연·기관 제출 불가] " if workflow and workflow.get("is_demo") else "") + "지역 분석 및 복지사업 검토 결과 보고서", "Title")
    document.add_paragraph(
        "저장된 지역 분석과 담당자의 복지사업 적합성 검토 결과를 정리합니다. 적합·보류·부적합 판단과 검토 근거를 기록합니다."
    )
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    metadata = [
        ("작성 부서 및 작성자", f"{department} / {author}"),
        ("작성일", datetime.now(ZoneInfo("Asia/Seoul")).strftime("%Y-%m-%d")),
        ("대상 지역 및 기간", f"{context['city']} / {context.get('district', '')} / {context['month']}"),
        ("자료 및 분석 버전", source_note),
    ]
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, str(value)
    document.add_heading("1 검토 요약", level=1)
    document.add_paragraph(
        f"해당 조건에서 확인할 변화 후보는 {len(rows)}건입니다. 아래 수치는 저장된 분석 결과이며, 자료 부족 또는 과거 이력 부족은 정상 판정과 구분하여 확인해야 합니다."
    )
    document.add_heading("2 관찰 사실과 탐지 근거", level=1)
    if rows:
        evidence = document.add_table(rows=1, cols=4)
        evidence.style = "Table Grid"
        for cell, name in zip(
            evidence.rows[0].cells,
            ["행정동 및 지표", "전월 변화율", "상대 변화량", "Robust Z"],
        ):
            cell.text = name
        for row in rows[:12]:
            cells = evidence.add_row().cells
            values = [
                f"{row['행정동명']}\n{row['metric_label']}",
                f"{row['change_pct']:+.1f}%" if row.get("change_pct") is not None else "자료 없음",
                f"{row['relative_change_pp']:+.1f}%p" if row.get("relative_change_pp") is not None else "자료 없음",
                f"{row['risk_robust_z']:.2f}" if row.get("risk_robust_z") is not None else "판단 보류",
            ]
            for cell, value in zip(cells, values):
                cell.text = value
        if len(rows) > 12:
            document.add_paragraph(
                f"총 {len(rows)}건 중 12건을 본문에 표시했습니다. 전체 결과는 데이터 관리 화면의 CSV를 참고하세요."
            )
    else:
        document.add_paragraph(
            "선택 조건에서 표시할 분석 후보가 없습니다. 후보 없음이 자료 충분 또는 사회적 고립 없음의 의미는 아닙니다."
        )
    document.add_heading("3 지역 맥락과 검토 사항", level=1)
    notes = list(
        dict.fromkeys(
            row.get("context_note", "") + " " + row.get("cluster_profile", "")
            for row in rows
        )
    )
    document.add_paragraph("\n".join(notes[:3]) or "연결된 지역 유형 자료가 없습니다.")
    document.add_paragraph(
        "활동 변화가 계절·지역 행사·수집 범위 변경에 따른 것인지 확인합니다. 지역 복지관·행정복지센터의 사회참여 및 생활 상담 자원을 검토하되 실제 사업·연락처·이용 조건은 별도 확인해야 합니다."
    )
    document.add_heading("4 담당자 검토 의견 및 후속 조치", level=1)
    document.add_paragraph(opinion.strip() or "담당자 검토 의견을 작성하세요.")
    document.add_heading("5 분석 기준과 자료 출처", level=1)
    document.add_paragraph(
        "상대 변화량은 동의 전월 변화율에서 같은 달 동 중앙값을 뺀 값입니다. 위험 방향 상대 변화 5%p 이상과 Robust Z-score 2.5 이상을 함께 적용합니다. 과거 변화 3회를 사용하며, MAD가 0이거나 이력이 부족하면 후보 판단을 보류합니다."
    )
    document.add_paragraph(source_note)
    if workflow:
        document.add_heading("6 분석 근거 확인 및 추가 질의", level=1)
        document.add_paragraph("근거 확인일: " + workflow.get("evidence_reviewed_at", "미확인"))
        for question in workflow.get("questions", []):
            document.add_paragraph("질문: " + question["question"])
            document.add_paragraph("답변: " + question.get("answer", ""))
        if not workflow.get("questions"):
            document.add_paragraph("추가 질의 없음 · 담당자가 분석 근거를 확인함")
        document.add_heading("7 복지사업 매칭 검토", level=1)
        for review in workflow.get("reviews", []):
            document.add_paragraph(f"사업: {review['name']} / 판단: {review['decision']}")
            document.add_paragraph("검토 근거: " + review["note"])
            document.add_paragraph("공식 출처: " + str(review.get("source_url") or "별도 확인"))
            snapshot = review.get("service_snapshot", {})
            for label, field in [("지원 대상", "target_text"), ("자격 조건", "eligibility_text"), ("거주 조건", "residency_text")]:
                document.add_paragraph(label + ": " + str(snapshot.get(field) or "별도 확인"))
        document.add_heading("8 보고 및 담당자 최종 확인", level=1)
        document.add_paragraph(f"작성 부서: {department} / 작성자: {author}")
        document.add_paragraph("보고 내용 및 후속 사항: " + opinion)
        document.add_paragraph("보고서 저장 후 대시보드 업무 완료로 처리됩니다. 기관 제출·공문 발송은 별도 확인합니다.")
    for table in document.tables:
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
    footer = section.footer.paragraphs[0]
    footer.text = "지역 분석 및 사업 검토 결과  |  기관 지정 양식 확정 전 내부 기본 양식"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
